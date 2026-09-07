# -*- coding: utf-8 -*-
"""
Stage 4 — ReportSpec 을 .pptx 로 렌더링

디자인은 전적으로 theme.py(스킬 포팅)의 헬퍼로 조립하며,
좌표·색상을 이 파일에서 즉흥적으로 새로 정하지 않는다.
"""

import re
from datetime import date
from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from . import theme as T
from . import formatters as fmt
from .blocks import ReportSpec, SlideSpec
from .emphasis import to_runs


def _day_ord(mmdd: str) -> int:
    m, d = mmdd.split('-')
    return date(2000, int(m), int(d)).toordinal()


class ReportRenderer:
    """ReportSpec → PPTX 파일"""

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(T.SLIDE_W)
        self.prs.slide_height = Inches(T.SLIDE_H)
        self._blank = self.prs.slide_layouts[6]
        self._tag = ''
        self._src_cache = {}       # 원본 pptx 경로 → Presentation (표 복제용)

    # ------------------------------------------------------------------

    #: 렌더 중 건너뛴 슬라이드 종류 (호출부가 Checklist 에 남길 수 있게 공개)
    skipped_kinds: List[str]
    #: 그리다 실패한 슬라이드 [(kind, 사유)] — 빈 장으로 남고 사실이 기록된다
    failed_slides: List[tuple]

    def render(self, spec: ReportSpec, out_path: str) -> str:
        self._tag = spec.campaign_tag
        self.skipped_kinds = []
        self.failed_slides = []
        for s in spec.slides:
            handler = getattr(self, f'_slide_{s.kind}', None)
            if handler is None:
                # 🔴 조용히 버리지 않는다.
                #
                # 예전에는 그냥 continue 했다. 지금은 모든 kind 에 핸들러가
                # 있지만, 새 블록을 추가하면서 핸들러를 빠뜨리면 그 슬라이드가
                # **아무 흔적 없이** 덱에서 사라진다. 장수만 줄어들 뿐이라
                # 아무도 눈치채지 못한다 (claude.md 3.1 조용한 소실 금지).
                self.skipped_kinds.append(s.kind)
                spec.warnings.append(
                    f"'{s.kind}' 슬라이드를 그릴 수 없어 건너뛰었어요 "
                    f'(렌더러에 해당 처리기가 없음)')
                continue

            # 🔴 한 장의 실패가 보고서 전체를 날리지 않게 한다.
            #
            # 핸들러들은 payload 를 `p['section']` 처럼 직접 인덱싱한다(65곳).
            # 블록 생성기가 키 하나를 빠뜨리면 KeyError 가 render() 밖으로
            # 터져 나가 **PPTX 가 아예 만들어지지 않는다.** AE 입장에선 몇 분
            # 걸린 작업이 마지막에 통째로 사라지는 셈이다. 시뮬레이션에서
            # 빈 payload 하나로 전체 렌더가 죽는 것을 확인했다.
            #
            # 이제 실패한 장은 비워 두고 사실을 기록한 뒤 계속 진행한다
            # (claude.md 3.3 파싱 에러 시 강제 종료 금지 · 3.1 조용한 소실 금지).
            slide = self.prs.slides.add_slide(self._blank)
            try:
                handler(slide, s.payload)
            except Exception as e:
                self.failed_slides.append((s.kind, f'{type(e).__name__}: {e}'))
                spec.warnings.append(
                    f"'{s.kind}' 슬라이드를 그리는 중 문제가 생겨 비워 뒀어요 "
                    f'— {type(e).__name__}')
                self._render_failure_notice(slide, s.kind)

        self._scrub_package()
        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return str(out)

    def _render_failure_notice(self, slide, kind: str) -> None:
        """
        그리지 못한 장을 빈 채로 두지 않고 왜 비었는지 한 줄 남긴다.

        아무 표시 없는 백지가 섞이면 AE 는 '원래 이런 장인가' 하고 넘어간다.
        보고서에 실린 채로 광고주에게 가는 것보다, 여기서 눈에 띄는 편이 낫다.
        """
        try:
            box = slide.shapes.add_textbox(
                Inches(1.0), Inches(3.0), Inches(11.3), Inches(1.2))
            tf = box.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            run = para.add_run()
            run.text = (f'이 장({kind})은 자동으로 그리지 못했어요. '
                        f'2단계에서 해당 항목을 확인해 주세요.')
            run.font.size = Pt(14)
            run.font.name = T.BODY_REG
        except Exception:
            # 안내 문구조차 실패하면 그냥 빈 장으로 둔다 — 여기서 또 터지면
            # 애초에 막으려던 '전체 렌더 실패'가 되돌아온다.
            pass

    # ------------------------------------------------------------------
    # 패키지 스크럽 — 절대 규칙을 파일 전체에 강제
    #
    # 기본 템플릿의 레이아웃/마스터/테마에는 bold 플래그와
    # 타 언어 스크립트용 폰트 목록(Angsana New, 맑은 고딕 등)이 남아 있다.
    # 우리 슬라이드는 쓰지 않지만, 기획자가 이 파일을 이어서 편집할 때
    # 그 서식을 상속받지 않도록 패키지 전체에서 제거한다.
    # ------------------------------------------------------------------

    _A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    _P = 'http://schemas.openxmlformats.org/presentationml/2006/main'

    def _is_original_table(self, el) -> bool:
        """복제한 원본 표(graphicFrame, cNvPr name='__ORIG_TBL__…')인지."""
        if el.tag != f'{{{self._P}}}graphicFrame':
            return False
        for cnv in el.iter(f'{{{self._P}}}cNvPr'):
            return cnv.get('name', '').startswith('__ORIG_TBL__')
        return False

    def _scrub_xml_root(self, root) -> None:
        """
        폰트·굵기 강제. 단, 복제한 원본 표 서브트리는 통째로 건너뛴다 —
        원본 표 디자인(색상·폰트 크기/위계)을 100% 보존하기 위함
        (사용자 지시 2026.09.06). lxml 프록시 id() 는 불안정하므로 트리
        구조를 직접 재귀하며 표식이 달린 graphicFrame 에서 가지치기한다.
        """
        A = self._A
        rpr_tags = {f'{{{A}}}rPr', f'{{{A}}}defRPr', f'{{{A}}}endParaRPr'}
        face_tags = {f'{{{A}}}latin', f'{{{A}}}ea', f'{{{A}}}cs',
                     f'{{{A}}}sym', f'{{{A}}}font', f'{{{A}}}buFont'}

        def recurse(el) -> None:
            if self._is_original_table(el):
                return                       # 복제 원본 표 — 전체 서브트리 보존
            if el.tag in rpr_tags:
                el.attrib.pop('b', None)
                el.attrib.pop('i', None)
            elif el.tag in face_tags:
                tf = el.get('typeface', '')
                if tf and not tf.startswith('+') and tf not in T.ALLOWED_FONTS:
                    el.set('typeface', T.BODY_REG)
            for child in el:
                recurse(child)

        recurse(root)

    def _scrub_package(self) -> None:
        import lxml.etree as etree
        from pptx.oxml.ns import qn

        for part in self.prs.part.package.iter_parts():
            name = str(part.partname)
            if not (name.startswith('/ppt/') and name.endswith('.xml')):
                continue
            element = getattr(part, '_element', None)
            if element is not None:
                # 슬라이드/레이아웃/마스터/차트 — 살아있는 트리 직접 수정
                self._scrub_xml_root(element)
                continue
            try:
                # 테마 등 blob 파트
                root = etree.fromstring(part.blob)
                # 테마 major/minor 대표 폰트는 의미에 맞게 지정
                for scheme, face in (('a:majorFont', T.HEAD_REG),
                                     ('a:minorFont', T.BODY_REG)):
                    for fs in root.iter(qn(scheme)):
                        for tag in ('a:latin', 'a:ea', 'a:cs'):
                            for e in fs.findall(qn(tag)):
                                e.set('typeface', face)
                self._scrub_xml_root(root)
                part._blob = etree.tostring(root, xml_declaration=True,
                                            encoding='UTF-8', standalone=True)
            except Exception:
                pass                     # 스크럽 실패가 저장을 막지 않도록

    def _new_footer(self, slide, sources: Optional[List[str]]):
        if sources:
            T.add_footnote(slide, sources)

    # ────────────────────────── [Checklist] (1페이지 강제)

    def _slide_checklist(self, slide, p):
        T.add_header(slide, self._tag, '[Checklist] 기획자 최종 검수 및 보정 항목')
        mode = '포스트바이 포함 (Full)' if p.get('mode') == 'full' else '포스트바이 없음 (Lite)'
        T.add_text(slide, 0.70, 1.10, 12.0, 0.3,
                   [(f'데이터 구성: {mode}', T.BODY_REG, 10, T.MUTED)])

        rows = p.get('rows', [])
        if not rows:
            T.add_text(slide, 0.70, 1.9, 12.0, 0.4,
                       [('확인이 필요한 예외 사항이 없습니다.', T.BODY_REG, 12, T.INK)])
        sev_color = {'error': T.NEG, 'warning': T.BLUE_MAIN, 'info': T.MUTED}
        y = 1.60
        for r in rows:
            c = sev_color.get(r['sev'], T.MUTED)
            runs = [(f"{r['mark']}  ", T.BODY_BOLD, 10, c),
                    (r['message'], T.BODY_REG, 10, T.BLACK)]
            if r.get('detail'):
                runs.append((f"  —  {r['detail']}", T.BODY_REG, 9, T.MUTED))
            T.add_text(slide, 0.70, y, 12.0, 0.34, runs)
            y += 0.36
        if p.get('overflow'):
            T.add_text(slide, 0.70, y + 0.05, 12.0, 0.3,
                       [(f"외 {p['overflow']}건 — 상세는 검증 스냅샷(JSON) 참조",
                         T.BODY_REG, 9, T.FOOT)])
        self._new_footer(slide, ['본 슬라이드는 기획자 더블 체크용이며 광고주 전달 전 확인 후 처리'])

    # ────────────────────────── 표지 / 목차 / 간지 / EOD

    def _slide_cover(self, slide, p):
        T.set_slide_background(slide, T.WHITE)
        title = p['title']
        size = 28 if len(title) <= 18 else 24
        T.add_text(slide, 1.0, 2.45, 11.3, 1.5,
                   [(title, T.HEAD_BOLD, size, T.BLACK)])
        T.add_text(slide, 1.0, 3.75, 11.3, 0.5,
                   [(p['subtitle'], T.BODY_REG, 16, T.INK)])
        T.add_text(slide, 1.0, 4.35, 11.3, 0.4,
                   [(p['date'], T.BODY_REG, 12, T.MUTED)])
        # 하단 그라데이션 띠 (표지에만 허용되는 장식)
        band = T.add_rect(slide, 0, T.SLIDE_H - 0.18, T.SLIDE_W, 0.18, T.BLUE_SOFT)
        try:
            band.fill.gradient()
            stops = band.fill.gradient_stops
            stops[0].color.rgb = T._rgb('5B8DEF')
            stops[1].color.rgb = T._rgb('A8C4F0')
            band.fill.gradient_angle = 0
        except Exception:
            pass                                    # 단색 폴백

    def _slide_toc(self, slide, p):
        T.set_slide_background(slide, T.DARK_NAVY)
        items = p.get('items', [])
        n = len(items)
        gap = 0.56 + 0.30
        y0 = (T.SLIDE_H - (n * gap - 0.30)) / 2
        for i, name in enumerate(items):
            T.add_text(slide, 0, y0 + i * gap, T.SLIDE_W, 0.56,
                       [(f'0{i + 1}. {name}', T.HEAD_BOLD, 20, T.WHITE)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def _slide_divider(self, slide, p):
        T.add_divider(slide, p['text'])

    def _slide_eod(self, slide, p):
        T.add_divider(slide, 'E.O.D', appendix=True, font_size=20)

    # ────────────────────────── 로드맵 (간트)

    def _slide_roadmap(self, slide, p):
        T.add_header(slide, self._tag, p['section'])
        if p.get('period_raw'):
            T.add_key_message(slide, [
                {'text': '캠페인 집행 기간  '},
                {'text': p['period_raw'], 'emph': True},
            ])

        bars = p['bars']
        d0 = min(_day_ord(b['start']) for b in bars) - 2
        d1 = max(_day_ord(b['end']) for b in bars) + 2
        span = max(d1 - d0, 1)

        gx0, gx1 = 2.75, 12.65
        gy0 = 2.75
        gy1 = 6.55
        row_h = min(0.42, (gy1 - gy0) / max(len(bars), 1))

        def X(day_ord: int) -> float:
            return gx0 + (gx1 - gx0) * (day_ord - d0) / span

        # 월 눈금
        cur = date.fromordinal(d0)
        seen_months = set()
        probe = d0
        while probe <= d1:
            dt = date.fromordinal(probe)
            if dt.day == 1 or (probe == d0 and dt.month not in seen_months):
                first = date(2000, dt.month, 1).toordinal()
                x = X(max(first, d0))
                seen_months.add(dt.month)
                T.add_rect(slide, x, gy0 - 0.05, 0.012, (gy1 - gy0) + 0.05,
                           T.LINE_SOFT)
                T.add_text(slide, x - 0.3, gy0 - 0.38, 0.8, 0.26,
                           [(f'{dt.month}월', T.BODY_REG, 9, T.MUTED)],
                           align=PP_ALIGN.CENTER)
            probe += 1

        for i, b in enumerate(bars):
            y = gy0 + i * row_h
            # 소재 축 로드맵처럼 라벨을 직접 지정하는 경우를 허용한다
            label = b.get('label') or b['media']
            if not b.get('label') and b.get('product') and b['product'] != label:
                label = f"{b['media']} · {b['product']}"
            if len(label) > 20:
                label = label[:19] + '…'
            T.add_text(slide, 0.70, y + row_h * 0.10, 1.95, row_h * 0.8,
                       [(label, T.BODY_REG, 9, T.BLACK)],
                       anchor=MSO_ANCHOR.MIDDLE)
            x_s, x_e = X(_day_ord(b['start'])), X(_day_ord(b['end']))
            bar_w = max(x_e - x_s, 0.06)
            T.add_rect(slide, x_s, y + row_h * 0.22, bar_w, row_h * 0.52,
                       T.BLUE_MAIN)
            period_txt = f"{b['start'].replace('-', '/')}~{b['end'].replace('-', '/')}"
            T.add_text(slide, min(x_e + 0.08, gx1 - 1.0), y + row_h * 0.14,
                       1.4, row_h * 0.7,
                       [(period_txt, T.BODY_REG, 8, T.MUTED)],
                       anchor=MSO_ANCHOR.MIDDLE)

        notes = list(p.get('sources') or [])
        if p.get('note'):
            notes.append(p['note'])
        self._new_footer(slide, notes)

    # ────────────────────────── 운영 요약

    def _slide_summary(self, slide, p):
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, p['key_parts'])

        cards = p['cards']
        cw, ch, gap = 2.70, 1.62, 0.35
        total_w = len(cards) * cw + (len(cards) - 1) * gap
        x0 = (T.SLIDE_W - total_w) / 2
        y0 = 3.45
        for i, (label, value) in enumerate(cards):
            x = x0 + i * (cw + gap)
            T.add_rect(slide, x, y0, cw, ch, T.WHITE, T.LINE_CARD, 0.75)
            T.add_rect(slide, x, y0, cw, 0.34, T.BLUE_MAIN)
            T.add_text(slide, x, y0, cw, 0.34,
                       [(label, T.BODY_REG, 11, T.WHITE)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            T.add_text(slide, x, y0 + 0.34, cw, ch - 0.34,
                       [(value, T.HEAD_MED, 20, T.INK)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self._new_footer(slide, p.get('sources'))

    # ────────────────────────── 일자별 추이

    def _slide_daily_trend(self, slide, p):
        T.add_header(slide, self._tag, p['section'])
        n_days = len(p['categories'])
        T.add_key_message(slide, [
            {'text': '캠페인 기간 '},
            {'text': f'{n_days}일', 'emph': True},
            {'text': ' 운영 · 운영 이벤트 '},
            {'text': f"{p['event_total']}건", 'emph': True},
            {'text': ' 기록'},
        ])
        T.add_block_label(slide, f"일자별 {p['metric_name']} 추이", 0.70, 2.62)
        T.add_line_chart(slide, p['categories'],
                         [{'name': p['metric_name'], 'values': p['values'],
                           'color': T.BLUE_SKY}],
                         x=0.70, y=2.95, w=7.9, h=3.75,
                         label_skip=max(1, n_days // 12))

        T.add_block_label(slide, '주요 운영 이벤트', 8.95, 2.62, 3.9)
        y = 2.98
        for ev in p['events']:
            T.add_text(slide, 8.95, y, 3.9, 0.4, [
                (f"{ev['date']}  ", T.BODY_BOLD, 9, T.BLUE_MAIN),
                (ev['note'][:34], T.BODY_REG, 9, T.INK),
            ])
            y += 0.42
        if p['event_total'] > len(p['events']):
            T.add_text(slide, 8.95, y, 3.9, 0.3,
                       [(f"외 {p['event_total'] - len(p['events'])}건",
                         T.BODY_REG, 9, T.FOOT)])
        self._new_footer(slide, p.get('sources'))

    # ────────────────────────── 매체별 표

    def _slide_media_table(self, slide, p):
        """매체별·소재별 집행 결과 표 (같은 표 스타일을 공유)"""
        T.add_header(slide, self._tag, p['section'])
        if p.get('headline_parts'):
            T.add_key_message(slide, p['headline_parts'])
        else:
            n_media = max(len(p['rows']) - 1, 0)
            T.add_key_message(slide, [
                {'text': 'Digital '},
                {'text': f'{n_media}개 매체', 'emph': True},
                {'text': ' 집행 결과 — 상세 지표 하기 표 기준'},
            ])
        if p.get('block_label'):
            T.add_block_label(slide, p['block_label'], 0.70, 2.62)
        n_rows = len(p['rows']) + 1
        row_h = 0.24 if n_rows <= 16 else 0.22
        T.add_styled_table(
            slide, p['header'], p['rows'],
            x=0.70, y=2.95, w=11.9,
            col_w=p.get('col_w') or [1.7, 1.5, 1.5, 1.3, 1.2, 0.9, 0.9, 1.0, 1.0],
            font_size=9 if n_rows <= 16 else 8,
            row_h=row_h, total_row=p.get('use_total_row', True))
        self._new_footer(slide, p.get('sources'))

    # ────────────────────────── 포스트바이 원본 표 (Deep-Dive 1:1 미러링)

    def _slide_postbuy_table(self, slide, p):
        """
        포스트바이 원본 표 1장을 미러링한다.

        1순위: 원본 pptx 의 표 도형을 그대로 복제 (스타일·색상·폰트/위계 100% 보존).
        복제 불가 시: 빌더가 넘긴 텍스트로 표를 재구성 (데이터 보존 폴백).
        """
        T.add_header(slide, self._tag, p['section'])
        if p.get('key_parts'):
            T.add_key_message(slide, p['key_parts'])

        cloned = False
        if p.get('src_pptx'):
            try:
                cloned = self._clone_original_table(
                    slide, p['src_pptx'], int(p.get('src_slide_no', 0)),
                    int(p.get('src_table_ord', 0)), top=2.72)
            except Exception:
                cloned = False        # 어떤 이유로든 실패하면 폴백

        if not cloned:
            T.add_styled_table(
                slide, p['header'], p['rows'],
                x=0.70, y=2.95, w=11.9,
                col_w=p.get('col_w'),
                font_size=p.get('font_size', 9),
                row_h=p.get('row_h', 0.24),
                total_row=p.get('use_total_row', False))
        self._new_footer(slide, p.get('sources'))

    # ------------------------------------------------------------------
    # 원본 표 도형 복제 — 원본 디자인을 그대로 가져온다
    # ------------------------------------------------------------------

    def _src_presentation(self, path: str):
        prs = self._src_cache.get(path)
        if prs is None:
            prs = Presentation(path)
            self._src_cache[path] = prs
        return prs

    @staticmethod
    def _collect_source_tables(shapes, out) -> None:
        """원본 슬라이드의 표 graphicFrame 을 파서와 동일한 순서로 수집."""
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                ReportRenderer._collect_source_tables(sh.shapes, out)
            elif getattr(sh, 'has_table', False):
                out.append(sh)

    def _clone_original_table(self, slide, src_pptx: str, slide_no: int,
                              table_ord: int, *, top: float) -> bool:
        """
        원본 pptx slide_no(1-base)의 table_ord 번째 표 도형을 복제해 삽입한다.

        · 원본 서식(셀 채움·테두리·병합·폰트·크기)을 그대로 유지
        · 본문 폭(11.9in)·높이 예산을 넘으면 비율을 지켜 축소
        · 스크럽 예외 표식(cNvPr name)을 달아 폰트 강제에서 제외
        """
        import copy
        from pptx.util import Inches, Emu
        from pptx.oxml.ns import qn

        prs = self._src_presentation(src_pptx)
        if slide_no < 1 or slide_no > len(prs.slides):
            return False
        tables = []
        self._collect_source_tables(prs.slides[slide_no - 1].shapes, tables)
        if table_ord < 0 or table_ord >= len(tables):
            return False

        gf = copy.deepcopy(tables[table_ord]._element)

        # 현재 표 크기 (xfrm 우선, 없으면 gridCol/tr 합)
        A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        xfrm = gf.find(qn('p:xfrm'))
        tbl = gf.find('.//' + qn('a:tbl'))
        if tbl is None:
            return False
        grid = tbl.find(qn('a:tblGrid'))
        cols = [int(c.get('w')) for c in grid.findall(qn('a:gridCol'))]
        trs = tbl.findall(qn('a:tr'))
        cur_w = sum(cols) if cols else None
        cur_h = sum(int(r.get('h', 0)) for r in trs) if trs else None
        ext = xfrm.find(qn('a:ext')) if xfrm is not None else None
        if ext is not None:
            cur_w = cur_w or int(ext.get('cx'))
            cur_h = cur_h or int(ext.get('cy'))
        if not cur_w or not cur_h:
            return False

        EMU = 914400
        avail_w = 11.9 * EMU
        avail_h = (T.FOOT_Y - top - 0.12) * EMU
        scale = min(1.0, avail_w / cur_w, avail_h / cur_h)

        if scale < 0.999:
            # 열폭·행높이·글자크기·ext 를 동일 비율로 축소 (원본 비율 유지)
            for c in grid.findall(qn('a:gridCol')):
                c.set('w', str(max(1, int(int(c.get('w')) * scale))))
            for r in trs:
                if r.get('h'):
                    r.set('h', str(max(1, int(int(r.get('h')) * scale))))
            for tag in ('a:rPr', 'a:defRPr', 'a:endParaRPr'):
                for rpr in tbl.iter(qn(tag)):
                    if rpr.get('sz'):
                        rpr.set('sz', str(max(100, int(int(rpr.get('sz')) * scale))))
            cur_w = int(cur_w * scale)
            cur_h = int(cur_h * scale)

        # 위치: 가로 중앙(본문 0.70~12.60), 세로 top
        left = int((0.70 * EMU) + max(0, (11.9 * EMU - cur_w) / 2))
        if xfrm is None:
            xfrm = gf.makeelement(qn('p:xfrm'), {})
            gf.insert(0, xfrm)
        for tag in ('a:off', 'a:ext'):
            for el in xfrm.findall(qn(tag)):
                xfrm.remove(el)
        off = xfrm.makeelement(qn('a:off'), {'x': str(int(left)),
                                             'y': str(int(top * EMU))})
        ext2 = xfrm.makeelement(qn('a:ext'), {'cx': str(int(cur_w)),
                                              'cy': str(int(cur_h))})
        xfrm.append(off)
        xfrm.append(ext2)

        # 스크럽 예외 표식 (원본 폰트/위계 보존)
        for cnv in gf.iter(qn('p:cNvPr')):
            cnv.set('name', '__ORIG_TBL__' + cnv.get('name', 'tbl'))
            break

        slide.shapes._spTree.append(gf)
        return True

    # ────────────────────────── 소재 카드

    def _slide_creative_cards(self, slide, p):
        """소재 카드 행 — 이미지(contain) + 캡션 + 지표"""
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, p['headline_parts'])
        T.add_block_label(slide, p['block_label'], 0.70, 2.62)

        cards = p['cards'][:5]
        n = max(len(cards), 1)
        gap = 0.35
        cw = (11.9 - gap * (n - 1)) / n
        ih = min(2.55, cw * 0.62)            # 이미지 상자 높이
        x0, y0 = 0.70, 3.00
        text_h = 1.05

        for i, card in enumerate(cards):
            cx = x0 + i * (cw + gap)
            T.add_rect(slide, cx, y0, cw, ih + text_h, T.WHITE, T.LINE_CARD, 0.75)

            placed = False
            if card.get('image'):
                try:
                    self._place_image(slide, card['image'],
                                      cx + 0.08, y0 + 0.08, cw - 0.16, ih - 0.16)
                    placed = True
                except Exception:
                    placed = False          # 손상·미지원 파일은 자리표시자로 대체
            if not placed:
                T.add_image_placeholder(slide, cx + 0.08, y0 + 0.08,
                                        cw - 0.16, ih - 0.16)

            ty = y0 + ih
            label = card['label']
            if len(label) > 18:
                label = label[:17] + '…'
            T.add_text(slide, cx + 0.10, ty + 0.04, cw - 0.20, 0.28,
                       [(label, T.BODY_BOLD, 10.5, T.INK)],
                       align=PP_ALIGN.CENTER)
            for k, line in enumerate(card['lines'][:3]):
                T.add_text(slide, cx + 0.10, ty + 0.32 + k * 0.21, cw - 0.20, 0.20,
                           to_runs(line, T.BODY_REG, T.BODY_BOLD, 9,
                                   T.MUTED, T.BLUE_MAIN),
                           align=PP_ALIGN.CENTER)

        self._new_footer(slide, p.get('sources'))

    @staticmethod
    def _place_image(slide, path: str, x: float, y: float,
                     w: float, h: float) -> None:
        """
        상자 안에 비율을 유지해 넣는다 (contain).

        design-spec 7장: 이미지 stretch 로 비율 깨기 금지.
        """
        pic = slide.shapes.add_picture(path, Inches(x), Inches(y),
                                       width=Inches(w))
        max_h = Inches(h)
        if pic.height > max_h:
            ratio = max_h / pic.height
            pic.height = int(max_h)
            pic.width = int(pic.width * ratio)
        pic.left = Inches(x) + int((Inches(w) - pic.width) / 2)
        pic.top = Inches(y) + int((Inches(h) - pic.height) / 2)

    # ────────────────────────── KPI

    def _slide_kpi(self, slide, p):
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        has_chart = bool(p['chart_items'])
        if has_chart:
            T.add_block_label(slide, '목표 대비 달성율 (제안=100 정규화)', 0.70, 2.62)
            T.add_achievement_chart(slide, p['chart_items'],
                                    x=0.70, y=3.05, w=4.6, h=3.55)
            tx, tw = 5.70, 6.95
        else:
            tx, tw = 0.70, 11.9
        T.add_block_label(slide, '매체 · 상품별 KPI 달성 현황', tx, 2.62)
        n_rows = len(p['table_rows']) + 1
        T.add_styled_table(
            slide, p['table_header'], p['table_rows'],
            x=tx, y=2.95, w=tw,
            col_w=[1.1, 1.3, 1.2, 1.2, 1.2, 0.9],
            font_size=8.5 if n_rows > 10 else 9,
            row_h=0.26 if n_rows <= 12 else 0.23)
        self._new_footer(slide, p.get('sources'))

    # ────────────────────────── 인용 (검색량/버즈)

    def _slide_quotes(self, slide, p):
        T.add_header(slide, self._tag, p['section'])
        quotes = p['quotes']
        y = 2.4
        for q in quotes:
            T.add_text(slide, 0.7, y, 11.9, 0.32,
                       [(q['label'], T.BODY_BOLD, 11, T.BLUE_MAIN)],
                       align=PP_ALIGN.CENTER)
            T.add_text(slide, 0.7, y + 0.36, 11.9, 0.8,
                       [(f'“ {q["text"]} ”', T.BODY_REG, 13, T.BLUE_EMPH)],
                       align=PP_ALIGN.CENTER)
            y += 1.35
        self._new_footer(slide, p.get('sources'))

    # ────────────────────────── Lesson Learned (자동 도출 3단 블록)

    # 밀도 후보 — (발견 pt, 제안 pt, 근거 pt, 근거 최대 건수,
    #              근거 1줄 강제, ▽ 인라인)
    # 느슨한 것부터 시도하고 전량이 들어가는 첫 후보를 채택한다.
    # 실제 보고서는 한 장에 3단 블록 4개를 담으므로 전량 게재를 우선한다.
    _DENSITY_STEPS = (
        (12.0, 11.0, 9.5, 3, False, False),
        (11.5, 10.5, 9.0, 3, False, False),
        (11.0, 10.0, 8.5, 3, False, False),
        (10.5, 9.5, 8.0, 3, False, False),
        (10.5, 9.5, 8.0, 2, False, True),
        (10.0, 9.0, 8.0, 2, True, True),
        (9.5, 8.5, 7.5, 1, True, True),
    )

    @staticmethod
    def _est_lines(text: str, width_in: float, size_pt: float) -> int:
        """
        한글 전각 글자는 폰트 크기와 거의 같은 폭을 차지한다.
        숫자·영문이 섞이면 더 짧아지므로 전각 1.0 / 반각 0.6 으로 환산한다.
        """
        span = sum(1.0 if ord(ch) > 0x2000 else 0.6 for ch in text)
        cpl = max(1.0, width_in * 72 / size_pt)
        return max(1, int(span / cpl) + 1)

    def _slide_placement_cards(self, slide, p):
        """매체별 게재 화면 — 실제 노출 캡처 + 집행 지표"""
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        cards = p['cards'][:4]
        n = max(len(cards), 1)
        gap = 0.30
        cw = (11.9 - gap * (n - 1)) / n
        ih = min(3.10, cw * 1.05)
        x0, y0 = 0.70, 2.90
        text_h = 0.78

        for i, card in enumerate(cards):
            cx = x0 + i * (cw + gap)
            T.add_rect(slide, cx, y0, cw, ih + text_h, T.WHITE, T.LINE_CARD, 0.75)

            placed = False
            if card.get('image'):
                try:
                    self._place_image(slide, card['image'],
                                      cx + 0.08, y0 + 0.08, cw - 0.16, ih - 0.16)
                    placed = True
                except Exception:
                    placed = False
            if not placed:
                T.add_image_placeholder(slide, cx + 0.08, y0 + 0.08,
                                        cw - 0.16, ih - 0.16, '게재 화면')

            ty = y0 + ih + 0.06
            T.add_text(slide, cx + 0.14, ty, cw - 0.28, 0.30,
                       [(card.get('media', ''), T.BODY_BOLD, 12.0, T.BLUE_MAIN)])
            if card.get('metric'):
                T.add_text(slide, cx + 0.14, ty + 0.32, cw - 0.28, 0.28,
                           [(card['metric'], T.BODY_REG, 10.5, T.MUTED)])

        self._new_footer(slide, p.get('sources') or [])

    # ────────────────────────── Part 1 캠페인 개요 (기획 의도)

    def _tip_box(self, slide, tip: str, y: float = 2.75) -> None:
        """
        소스 미확보 시 까는 AE 작성 가이드.

        슬라이드를 지우지 않는 이유: 목차·페이지네이션이 깨지고, 무엇보다
        '기획 의도가 비었다'는 사실 자체가 보고서에서 사라진다.
        """
        x, w = 0.70, 11.9
        lines = [ln.strip() for ln in (tip or '').split('\n') if ln.strip()]
        lh = 0.30
        h = 0.34 + lh * len(lines) + 0.20
        T.add_rect(slide, x, y, w, h, T.HILITE, T.LINE_SOFT, 1.0)
        iy = y + 0.20
        for idx, ln in enumerate(lines):
            bold = idx <= 1                       # 머리 두 줄은 강조
            T.add_text(slide, x + 0.34, iy, w - 0.68, lh,
                       [(ln, T.BODY_BOLD if bold else T.BODY_REG,
                         12.5 if idx == 0 else 11.0,
                         T.INK if bold else T.MUTED)])
            iy += lh

    def _slide_ov_goal(self, slide, p):
        """캠페인 목표 — 당면 과제 · 핵심 타겟 · 메인 카피"""
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        if p.get('tip'):
            self._tip_box(slide, p['tip'])
            self._new_footer(slide, p.get('sources') or [])
            return

        x, w = 0.70, 11.9
        y, gap = 2.72, 0.16
        limit = T.FOOT_Y - 0.10
        for label, text in p['rows']:
            body = re.sub(r'\s+', ' ', text).strip()
            n = self._est_lines(body, w - 1.90, 12.0)
            bh = max(0.62, 0.26 + 0.225 * n)
            if y + bh > limit:
                break
            T.add_rect(slide, x, y, w, bh, T.WHITE, T.LINE_CARD, 0.75)
            T.add_rect(slide, x, y, 0.055, bh, T.BLUE_MAIN, None, 0)
            T.add_text(slide, x + 0.26, y + 0.13, 1.45, 0.30,
                       [(label, T.BODY_BOLD, 12.0, T.BLUE_MAIN)])
            T.add_text(slide, x + 1.78, y + 0.13, w - 2.10, bh - 0.26,
                       to_runs(body, T.BODY_REG, T.BODY_BOLD, 12.0,
                               T.INK, T.BLUE_EMPH))
            y += bh + gap
        self._new_footer(slide, p.get('sources') or [])

    def _slide_ov_strategy(self, slide, p):
        """캠페인 전략 — Phase 구성 + 핵심 채널"""
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        if p.get('tip'):
            self._tip_box(slide, p['tip'])
            self._new_footer(slide, p.get('sources') or [])
            return

        x, w = 0.70, 11.9
        y = 2.72
        phases = (p.get('phases') or [])[:4]

        if phases:
            cols = len(phases)
            cw = (w - 0.18 * (cols - 1)) / cols
            ch = 2.30
            for idx, ph in enumerate(phases):
                cx = x + idx * (cw + 0.18)
                T.add_rect(slide, cx, y, cw, ch, T.WHITE, T.LINE_CARD, 0.75)
                T.add_rect(slide, cx, y, cw, 0.46, T.TBL_HEAD, None, 0)
                name = f"Phase{idx + 1}. {ph.get('name', '')}".strip()
                T.add_text(slide, cx + 0.18, y + 0.09, cw - 0.36, 0.30,
                           [(name, T.BODY_BOLD, 12.0, T.BLUE_MAIN)])
                iy = y + 0.58
                if ph.get('period'):
                    T.add_text(slide, cx + 0.18, iy, cw - 0.36, 0.26,
                               [(ph['period'], T.BODY_BOLD, 10.5, T.MUTED)])
                    iy += 0.28
                purpose = re.sub(r'\s+', ' ', ph.get('purpose') or '').strip()
                if purpose:
                    T.add_text(slide, cx + 0.18, iy, cw - 0.36, ch - (iy - y) - 0.16,
                               to_runs(purpose, T.BODY_REG, T.BODY_BOLD, 10.5,
                                       T.INK, T.BLUE_EMPH))
            y += ch + 0.24

        channels = p.get('channels') or []
        if channels:
            T.add_text(slide, x, y, 3.0, 0.28,
                       [('핵심 채널', T.BODY_BOLD, 11.5, T.BLUE_MAIN)])
            y += 0.34
            cx, chip_h = x, 0.34
            for name in channels[:12]:
                cwid = min(2.6, 0.30 + 0.115 * len(name) + 0.20)
                if cx + cwid > x + w:
                    cx = x
                    y += chip_h + 0.10
                T.add_rect(slide, cx, y, cwid, chip_h, T.BG_CARD, T.LINE_CARD, 0.6)
                T.add_text(slide, cx + 0.12, y + 0.05, cwid - 0.24, 0.24,
                           [(name, T.BODY_REG, 10.5, T.INK)])
                cx += cwid + 0.10

        self._new_footer(slide, p.get('sources') or [])

    def _slide_ov_roadmap(self, slide, p):
        """캠페인 로드맵 — Phase 순차 타임라인 (기획 의도 기준)"""
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        if p.get('tip'):
            self._tip_box(slide, p['tip'])
            self._new_footer(slide, p.get('sources') or [])
            return

        x, w = 0.70, 11.9
        y = 2.80
        phases = (p.get('phases') or [])[:5]
        n = len(phases)

        # 시간축
        T.add_rect(slide, x, y + 0.30, w, 0.03, T.LINE_SOFT, None, 0)
        seg = w / max(1, n)
        for idx, ph in enumerate(phases):
            sx = x + idx * seg
            T.add_rect(slide, sx + 0.04, y + 0.10, seg - 0.16, 0.44,
                       T.BLUE_MAIN if idx == 0 else T.BLUE_SOFT, None, 0)
            label = f"Phase{idx + 1}. {ph.get('name', '')}".strip()
            T.add_text(slide, sx + 0.18, y + 0.16, seg - 0.40, 0.30,
                       [(label, T.BODY_BOLD, 11.5, T.WHITE)])

        by = y + 0.74
        for idx, ph in enumerate(phases):
            sx = x + idx * seg
            T.add_rect(slide, sx + 0.04, by, seg - 0.16, 1.90,
                       T.WHITE, T.LINE_CARD, 0.75)
            iy = by + 0.16
            if ph.get('period'):
                T.add_text(slide, sx + 0.20, iy, seg - 0.44, 0.28,
                           [(ph['period'], T.BODY_BOLD, 11.0, T.BLUE_MAIN)])
                iy += 0.32
            purpose = re.sub(r'\s+', ' ', ph.get('purpose') or '').strip()
            if purpose:
                T.add_text(slide, sx + 0.20, iy, seg - 0.44, 1.90 - (iy - by) - 0.16,
                           to_runs(purpose, T.BODY_REG, T.BODY_BOLD, 10.5,
                                   T.INK, T.BLUE_EMPH))

        if p.get('period'):
            T.add_text(slide, x, by + 2.04, w, 0.28,
                       [('전체 기간 ', T.BODY_REG, 10.5, T.MUTED),
                        (p['period'], T.BODY_BOLD, 10.5, T.INK)])

        self._new_footer(slide, p.get('sources') or [])

    def _slide_insight(self, slide, p):
        """Lesson Learned (자동 도출) — 발견 → 제안 → 근거 3단 블록"""
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        x, w = 0.70, 11.9
        tw = w - 0.48                      # 블록 내부 텍스트 폭
        y0, gap = 2.58, 0.12
        limit = T.FOOT_Y - 0.10
        blocks = p['blocks']

        def line_h(size_pt: float) -> float:
            return size_pt * 1.34 / 72     # 줄 간격 포함 한 줄 높이

        def plan(step):
            """한 밀도 후보로 블록별 배치를 계산한다"""
            f_pt, r_pt, e_pt, e_max, one_line, inline = step
            pad = 0.12 if inline else 0.14
            arrow_h = 0.05 if inline else 0.19
            out = []
            for blk in blocks:
                finding = re.sub(r'\s+', ' ', blk['finding']).strip()
                rec = re.sub(r'\s+', ' ', blk['recommendation']).strip()
                evid = [re.sub(r'\s+', ' ', e).strip()
                        for e in blk['evidence']][:e_max]
                if one_line:
                    # '- ' 접두를 포함해 한 줄에 들어갈 때까지 줄인다.
                    # 글자수로 자르면 전각·반각 폭과 접두 폭이 빠져 한 줄을 넘긴다.
                    shortened = []
                    for e in evid:
                        while (len(e) > 12
                               and self._est_lines(f'- {e}', tw, e_pt) > 1):
                            e = e[:-4].rstrip(' +,(/·') + '…'
                        shortened.append(e)
                    evid = shortened

                # [평가] 3단 논법의 2단 — 수치의 기획적 해석
                ctx = re.sub(r'\s+', ' ', (blk.get('context') or '')).strip()
                ctx_prefix = '평가 · '

                rec_prefix = '▽  제안 · ' if inline else '제안 · '
                f_lines = self._est_lines(finding, tw, f_pt)
                c_lines = (self._est_lines(ctx_prefix + ctx, tw, r_pt)
                           if ctx else 0)
                r_lines = self._est_lines(rec_prefix + rec, tw, r_pt)
                e_lines = [self._est_lines(f'- {e}', tw, e_pt) for e in evid]
                fh = line_h(f_pt) * f_lines
                ch = line_h(r_pt) * c_lines
                rh = line_h(r_pt) * r_lines
                eh = sum(line_h(e_pt) * n for n in e_lines)
                out.append({
                    'finding': finding, 'ctx': ctx, 'rec': rec, 'evid': evid,
                    'ctx_prefix': ctx_prefix,
                    'rec_prefix': rec_prefix, 'arrow_h': arrow_h, 'pad': pad,
                    'fh': fh, 'ch': ch, 'rh': rh, 'e_lines': e_lines,
                    'h': (pad + fh + (ch + 0.03 if ctx else 0) + arrow_h + rh
                          + (0.04 + eh if evid else 0) + pad),
                })
            return out

        step, laid = self._DENSITY_STEPS[-1], None
        for candidate in self._DENSITY_STEPS:
            trial = plan(candidate)
            total = sum(b['h'] for b in trial) + gap * (len(trial) - 1)
            if y0 + total <= limit:
                step, laid = candidate, trial
                break
        if laid is None:
            laid = plan(step)
        f_pt, r_pt, e_pt = step[0], step[1], step[2]
        inline = step[5]

        y, shown = y0, 0
        for blk in laid:
            if shown and y + blk['h'] > limit:
                break
            T.add_rect(slide, x, y, w, blk['h'], T.WHITE, T.LINE_CARD, 0.75)
            iy = y + blk['pad']

            # [발견] 수치·지표명만 강조 런으로 분리
            T.add_text(slide, x + 0.24, iy, tw, blk['fh'],
                       to_runs(blk['finding'], T.BODY_REG, T.BODY_BOLD, f_pt,
                               T.INK, T.BLUE_MAIN))
            iy += blk['fh']

            # [평가] 기획적 해석 — 발견과 제안 사이의 논리 연결
            if blk['ctx']:
                T.add_text(slide, x + 0.24, iy, tw, blk['ch'],
                           [(blk['ctx_prefix'], T.BODY_BOLD, r_pt, T.MUTED)] +
                           to_runs(blk['ctx'], T.BODY_REG, T.BODY_BOLD, r_pt,
                                   T.INK, T.BLUE_SOFT))
                iy += blk['ch'] + 0.03

            # ▽ 전환 표시 — 블록이 많을 때는 제안 줄에 인라인으로 붙인다
            if not inline:
                T.add_text(slide, x + 0.24, iy, 0.4, 0.17,
                           [('▽', T.BODY_REG, e_pt, T.BLUE_MAIN)])
            iy += blk['arrow_h']

            # [제안]
            T.add_text(slide, x + 0.24, iy, tw, blk['rh'],
                       [(blk['rec_prefix'], T.BODY_BOLD, r_pt, T.BLUE_EMPH)] +
                       to_runs(blk['rec'], T.BODY_REG, T.BODY_BOLD, r_pt,
                               T.INK, T.BLUE_EMPH))
            iy += blk['rh'] + 0.04

            # [근거]
            for e, n in zip(blk['evid'], blk['e_lines']):
                T.add_text(slide, x + 0.24, iy, tw, line_h(e_pt) * n,
                           [('- ', T.BODY_REG, e_pt, T.MUTED)] +
                           to_runs(e, T.BODY_REG, T.BODY_REG, e_pt,
                                   T.MUTED, T.BLUE_SOFT))
                iy += line_h(e_pt) * n

            y += blk['h'] + gap
            shown += 1

        notes = list(p.get('sources') or [])
        if shown < len(blocks):
            notes.append(f'지면 제약으로 {len(blocks) - shown}건 미게재 '
                         f'— Checklist 참조')
        self._new_footer(slide, notes)

    # ────────────────────────── 차기 캠페인 전략 방향

    def _slide_strategy(self, slide, p):
        T.add_header(slide, self._tag, p['section'])
        T.add_key_message(slide, [{'text': p['headline']}])

        cards = p['cards'][:4]
        cols = 2 if len(cards) > 1 else 1
        rows = (len(cards) + cols - 1) // cols
        gap = 0.40
        cw = (11.9 - gap * (cols - 1)) / cols
        ch = min(1.95, (T.FOOT_Y - 0.20 - 2.62 - gap * (rows - 1)) / rows)
        x0, y0 = 0.70, 2.62

        for i, card in enumerate(cards):
            cx = x0 + (i % cols) * (cw + gap)
            cy = y0 + (i // cols) * (ch + gap)
            T.add_rect(slide, cx, cy, cw, ch, T.WHITE, T.LINE_CARD, 0.75)
            T.add_rect(slide, cx, cy, cw, 0.34, T.BLUE_MAIN)
            T.add_text(slide, cx, cy, cw, 0.34,
                       [(card['area'], T.BODY_REG, 11, T.WHITE)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

            direction = re.sub(r'\s+', ' ', card['direction']).strip()
            basis = re.sub(r'\s+', ' ', card['basis']).strip()
            T.add_text(slide, cx + 0.22, cy + 0.48, cw - 0.44, 0.66,
                       to_runs(direction, T.BODY_REG, T.BODY_BOLD, 11,
                               T.INK, T.BLUE_MAIN))
            T.add_text(slide, cx + 0.22, cy + 1.16, cw - 0.44, ch - 1.30,
                       [('근거 · ', T.BODY_BOLD, 9, T.MUTED)] +
                       to_runs(basis, T.BODY_REG, T.BODY_REG, 9,
                               T.MUTED, T.BLUE_SOFT))

        self._new_footer(slide, p.get('sources'))

    # ────────────────────────── Lesson Learned (포스트바이 원문)

    def _slide_lesson(self, slide, p):
        """Lesson Learned — [헤드라인 + 근거] 분석형 카드 (원문 위계 유지)"""
        T.add_header(slide, self._tag, p['section'])
        if p.get('headline_parts'):
            T.add_key_message(slide, p['headline_parts'])

        cards = p.get('cards') or []
        x, w = 0.70, 11.9
        tw = w - 0.52
        y = 2.66
        limit = T.FOOT_Y - 0.10
        gap = 0.14

        def line_h(sz: float) -> float:
            return sz * 1.34 / 72

        shown = 0
        for card in cards:
            finding = re.sub(r'\s+', ' ', card['finding']).strip()
            evid = [re.sub(r'\s+', ' ', e).strip()
                    for e in card.get('evidence', []) if e.strip()]
            f_pt, e_pt = 11.5, 9.5
            f_lines = self._est_lines(finding, tw, f_pt)
            e_lines = [self._est_lines('- ' + e, tw, e_pt) for e in evid]
            fh = line_h(f_pt) * f_lines
            eh = sum(line_h(e_pt) * n for n in e_lines)
            pad = 0.13
            ch = pad + fh + (0.06 + eh if evid else 0) + pad
            if shown and y + ch > limit:
                break
            T.add_rect(slide, x, y, w, ch, T.WHITE, T.LINE_CARD, 0.75)
            T.add_rect(slide, x, y, 0.09, ch, T.BLUE_MAIN)      # 좌측 액센트
            iy = y + pad
            T.add_text(slide, x + 0.28, iy, tw, fh,
                       to_runs(finding, T.BODY_BOLD, T.BODY_BOLD, f_pt,
                               T.INK, T.BLUE_MAIN))
            iy += fh + 0.06
            for e, n in zip(evid, e_lines):
                T.add_text(slide, x + 0.30, iy, tw - 0.04, line_h(e_pt) * n,
                           [('- ', T.BODY_REG, e_pt, T.MUTED)] +
                           to_runs(e, T.BODY_REG, T.BODY_REG, e_pt,
                                   T.MUTED, T.BLUE_SOFT))
                iy += line_h(e_pt) * n
            y += ch + gap
            shown += 1

        notes = list(p.get('sources') or [])
        if shown < len(cards):
            notes.append(f'지면 관계상 {len(cards) - shown}건은 다음 장에 이어짐')
        self._new_footer(slide, notes)
