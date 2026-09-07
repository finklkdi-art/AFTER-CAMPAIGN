# -*- coding: utf-8 -*-
"""
게재 보고(placement_report) — 실제 광고 게재 화면 이미지 추출

'게재 보고'는 매체별로 광고가 실제 노출된 화면을 캡처해 광고주에게 보고하는
문서다. 결과보고서에 실을 **게재 화면 이미지의 원천**이며, 다른 문서에는 없는
자산이다.

기존 creative_assets 는 캠페인 폴더에 흩어진 이미지 '파일'만 다뤘다. 게재보고는
이미지가 문서 안에 박혀 있어 따로 뽑아야 한다.

  PPTX : 슬라이드의 PICTURE 도형에서 blob 추출 + 같은 슬라이드 텍스트를 라벨로
  PDF  : 페이지 자체가 캡처 레이아웃이므로 페이지를 이미지로 렌더

추출 후 라벨(슬라이드 텍스트)을 미디어믹스·데일리리포트·포스트바이의 매체명과
대조해 어느 매체의 게재 화면인지 확정한다. 확정되지 않은 것은 버리지 않고
matched=False 로 남겨 Checklist 로 흘려보낸다 (조용한 소실 금지).
"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# 아이콘·로고 같은 장식 이미지를 걸러내는 하한.
# 용량으로 거르면 압축이 잘 된 스크린샷(단색 UI)이 통째로 날아간다 — 실측 5KB.
# 1차 신호는 '슬라이드에서 차지하는 크기'로 두고, 용량은 명백한 아이콘만 쳐낸다.
MIN_BYTES = 3 * 1024          # 이보다 작으면 아이콘으로 본다
MIN_BYTES_NO_SIZE = 8 * 1024  # 표시 크기를 못 구했을 때의 대체 하한
MIN_SIDE = 220                # 표시 크기(px) — 가로·세로 모두 이보다 작으면 장식
MAX_SHOTS = 40

_SEP = re.compile(r'[\s\-_.·/|,()\[\]]+')


@dataclass
class PlacementShot:
    """게재 화면 한 장"""
    image_path: str = ''          # 프로젝트 루트 기준 상대경로
    label: str = ''               # 캡처 주변 텍스트 (매체 판별 단서)
    media: str = ''               # 교차검증으로 확정된 매체명
    matched: bool = False
    source_file: str = ''
    slide_no: int = 0
    width: int = 0
    height: int = 0

    def to_dict(self) -> dict:
        return {
            'image_path': self.image_path, 'label': self.label,
            'media': self.media, 'matched': self.matched,
            'source_file': self.source_file, 'slide_no': self.slide_no,
            'width': self.width, 'height': self.height,
        }

    @classmethod
    def from_dict(cls, d: dict) -> 'PlacementShot':
        return cls(**{k: d.get(k, getattr(cls, k, '')) for k in
                      ('image_path', 'label', 'media', 'matched',
                       'source_file', 'slide_no', 'width', 'height')})


def _tokens(text: str) -> List[str]:
    return [t for t in _SEP.split((text or '').lower()) if len(t) >= 2]


# ────────────────────────── 추출

def _shot_dir(out_dir: Path, stem: str) -> Path:
    safe = re.sub(r'[^0-9A-Za-z가-힣]+', '_', stem)[:40] or 'placement'
    d = out_dir / safe
    d.mkdir(parents=True, exist_ok=True)
    return d


def _from_pptx(path: Path, out_dir: Path,
               project_root: Optional[Path]) -> List[PlacementShot]:
    from pptx import Presentation
    from pptx.util import Emu

    shots: List[PlacementShot] = []
    prs = Presentation(str(path))
    dest = _shot_dir(out_dir, path.stem)

    def texts_of(shapes) -> List[str]:
        out = []
        for sh in shapes:
            try:
                if sh.shape_type is not None and int(sh.shape_type) == 6 \
                        and hasattr(sh, 'shapes'):
                    out.extend(texts_of(sh.shapes))
                    continue
            except Exception:
                pass
            try:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    out.append(' '.join(sh.text_frame.text.split()))
            except Exception:
                pass
        return out

    def pics_of(shapes) -> List[Any]:
        out = []
        for sh in shapes:
            try:
                if sh.shape_type is not None and int(sh.shape_type) == 6 \
                        and hasattr(sh, 'shapes'):
                    out.extend(pics_of(sh.shapes))
                    continue
            except Exception:
                pass
            if getattr(sh, 'shape_type', None) is not None and \
                    str(sh.shape_type).startswith('PICTURE'):
                out.append(sh)
        return out

    for si, slide in enumerate(prs.slides, 1):
        label = ' / '.join(texts_of(slide.shapes))[:160]
        for pi, pic in enumerate(pics_of(slide.shapes), 1):
            if len(shots) >= MAX_SHOTS:
                return shots
            try:
                blob = pic.image.blob
                ext = (pic.image.ext or 'png').lower()
            except Exception:
                continue
            if len(blob) < MIN_BYTES:
                continue
            try:
                w = int(Emu(pic.width).inches * 96) if pic.width else 0
                h = int(Emu(pic.height).inches * 96) if pic.height else 0
            except Exception:
                w = h = 0
            if w and h:
                # 표시 크기가 1차 신호 — 가로·세로 모두 작으면 장식 이미지
                if w < MIN_SIDE and h < MIN_SIDE:
                    continue
            elif len(blob) < MIN_BYTES_NO_SIZE:
                continue

            fp = dest / f's{si:03d}_{pi:02d}.{ext}'
            try:
                fp.write_bytes(blob)
            except Exception:
                continue
            shots.append(PlacementShot(
                image_path=_rel(fp, project_root), label=label,
                source_file=path.name, slide_no=si, width=w, height=h))
    return shots


def _from_pdf(path: Path, out_dir: Path,
              project_root: Optional[Path]) -> List[PlacementShot]:
    """PDF 는 페이지 자체가 캡처 레이아웃이므로 페이지를 렌더한다"""
    import pdfplumber
    import pypdfium2 as pdfium

    shots: List[PlacementShot] = []
    dest = _shot_dir(out_dir, path.stem)

    labels: Dict[int, str] = {}
    try:
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                try:
                    t = (page.extract_text() or '').strip()
                except Exception:
                    t = ''
                labels[i] = ' / '.join(x.strip() for x in t.split('\n')
                                       if x.strip())[:160]
    except Exception:
        pass

    doc = pdfium.PdfDocument(str(path))
    for i in range(len(doc)):
        if len(shots) >= MAX_SHOTS:
            break
        try:
            page = doc[i]
            w, h = page.get_size()
            img = page.render(scale=1400 / w).to_pil().convert('RGB')
        except Exception:
            continue
        fp = dest / f'p{i + 1:03d}.png'
        try:
            img.save(fp, optimize=True)
        except Exception:
            continue
        shots.append(PlacementShot(
            image_path=_rel(fp, project_root), label=labels.get(i + 1, ''),
            source_file=path.name, slide_no=i + 1,
            width=img.width, height=img.height))
    return shots


def _rel(p: Path, project_root: Optional[Path]) -> str:
    """절대경로를 데이터에 남기지 않는다 (claude.md 1.1)"""
    try:
        if project_root:
            return str(p.resolve().relative_to(Path(project_root).resolve()))
    except Exception:
        pass
    return str(p)


def extract(documents: List[Any], out_dir: Any,
            project_root: Optional[Any] = None) -> List[PlacementShot]:
    """
    게재보고 문서에서 게재 화면 이미지를 뽑는다.

    Args:
        documents: SourceDocument 목록 (role == 'placement_report' 만 사용)
        out_dir:   이미지를 저장할 폴더
        project_root: 상대경로 기준
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    root = Path(project_root) if project_root else None

    shots: List[PlacementShot] = []
    for doc in documents or []:
        if getattr(doc, 'role', '') != 'placement_report':
            continue
        raw = getattr(doc, 'file_path', '') or ''
        p = Path(raw)
        if not p.is_absolute() and root:
            p = root / raw
        if not p.exists():
            continue
        try:
            if p.suffix.lower() == '.pptx':
                shots.extend(_from_pptx(p, out, root))
            elif p.suffix.lower() == '.pdf':
                shots.extend(_from_pdf(p, out, root))
        except Exception:
            continue
    return shots


# ────────────────────────── 매체 교차검증

def match_media(shots: List[PlacementShot],
                media_names: List[str]) -> Tuple[int, int]:
    """
    라벨을 실제 집행 매체명과 대조해 매체를 확정한다.

    미디어믹스·데일리리포트·포스트바이에서 모은 매체명만 정답으로 인정한다.
    문서에 적힌 이름을 그대로 믿지 않고, 실제 집행된 매체와 교차 검증하는 것이
    목적이다.

    Returns:
        (확정 건수, 미확정 건수)
    """
    names = [n for n in (media_names or []) if n and len(n) >= 2]
    ok = 0
    for shot in shots:
        hay = f'{shot.label} {Path(shot.image_path).name}'.lower()
        hay_tokens = set(_tokens(hay))
        best, best_score = '', 0
        for name in names:
            low = name.lower()
            score = 0
            if low in hay:
                score = 3 + len(low) // 4
            else:
                toks = [t for t in _tokens(name) if len(t) >= 2]
                if toks and all(t in hay_tokens for t in toks):
                    score = 2
                elif any(t in hay_tokens for t in toks):
                    score = 1
            if score > best_score:
                best, best_score = name, score
        if best_score >= 2:
            shot.media = best
            shot.matched = True
            ok += 1
    return ok, len(shots) - ok
