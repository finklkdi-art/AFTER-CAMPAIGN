# -*- coding: utf-8 -*-
"""
PostbuyParser — 포스트바이 리포트(PPTX)에서 정제된 보고 데이터를 추출

포스트바이는 표준 템플릿을 따른다(실측 2건 확인):
  표지 → Campaign Summary → KPI 달성 현황 → 목표 대비 달성율 → Scheduling
       → 검색어/검색량/버즈량 → 매체별 집행 결과 → 도달·빈도
       → Adobe Analytics → Lesson Learned → E.O.D

이 파서가 확보하는 것 중 데일리리포트로 대체 불가한 항목:
  KPI 확정 목표치 / 도달·빈도 / 검색량·버즈량 / Adobe Analytics / Lesson Learned
"""

import re
from pathlib import Path
from typing import Dict, List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from models.campaign_data import (
    KpiTarget, MediaPerformance, MetricSet, PostbuyResult, PostbuySection,
    Provenance, ScheduleEvent,
)
from . import sheet_utils as su


# 슬라이드 분류 규칙 — 제목/본문 텍스트에 대한 정규식 (앞선 것 우선)
_SLIDE_KINDS = (
    ('summary',    r'campaign\s*summary'),
    ('kpi',        r'kpi\s*달성|목표\s*대비\s*달성|달성\s*현황'),
    ('scheduling', r'scheduling|스케줄'),
    ('search',     r'검색어\s*추이|검색량\s*트렌드|관련\s*검색어'),
    ('buzz',       r'버즈량'),
    ('reach',      r'도달\s*및\s*빈도|reach'),
    ('analytics',  r'adobe\s*analytics'),
    ('lesson',     r'lesson\s*(&|and)?\s*learned|제언'),
    ('media',      r'매체\s*집행\s*결과|집행\s*결과|소재별|타겟팅별|지역별'),
)

# 이벤트 주석: "4/25", "6/1 ~ 6/30" 등으로 시작하는 짧은 텍스트
_EVENT_DATE_RE = re.compile(r'^\s*(\d{1,2})\s*/\s*(\d{1,2})\b')

# 기간 표기 "(기간 : 25.04.27~07.26)"
_PERIOD_RE = re.compile(r'기간\s*[:：]\s*([^)\n]+)')


class PostbuyParser:
    """포스트바이 PPTX 파서"""

    @staticmethod
    def parse(file_path: str) -> PostbuyResult:
        """
        Args:
            file_path: 포스트바이 pptx 경로

        Returns:
            PostbuyResult (실패해도 예외 없이 warnings 에 기록)
        """
        path = Path(file_path)
        result = PostbuyResult(source_file=path.name)

        try:
            prs = Presentation(str(path))
        except Exception as e:
            from utils.drm_check import detect_drm
            result.warnings.append(detect_drm(str(path))
                                   or f'PPTX 를 열 수 없음: {e}')
            return result

        for idx, slide in enumerate(prs.slides, 1):
            texts, tables = PostbuyParser._read_slide(slide)
            section = PostbuySection(
                slide_no=idx,
                kind=PostbuyParser._classify(texts),
                title=texts[0] if texts else '',
                headlines=[t for t in texts if len(t) > 12][:6],
                tables=tables,
                notes=[t for t in texts if t.startswith('*')],
            )
            result.sections.append(section)

            if idx == 1:
                PostbuyParser._read_cover(texts, result)
            if section.kind == 'summary':
                PostbuyParser._read_summary(section, path.name, result)
            elif section.kind == 'kpi':
                PostbuyParser._read_kpi(section, path.name, result)
            elif section.kind == 'scheduling':
                PostbuyParser._read_events(section, path.name, result)
            elif section.kind == 'lesson':
                # 간지 슬라이드는 제목만 있으므로 본문이 있는 것만 수집.
                # 표 안에 서술된 경우도 있어(무풍콤보 p20) 표 셀도 함께 훑는다.
                result.lesson_learned.extend(t for t in texts if len(t) > 20)
                for table in tables:
                    for row in table:
                        result.lesson_learned.extend(
                            su.norm(c) for c in row if len(su.norm(c)) > 20)

        if not result.campaign_summary:
            result.warnings.append('Campaign Summary 표를 찾지 못함')
        if not result.kpi_targets:
            result.warnings.append('KPI 달성 현황 표를 찾지 못함 — 확정 목표치 미확보')

        return result

    # ------------------------------------------------------------------

    @staticmethod
    def _read_slide(slide) -> tuple:
        """
        슬라이드의 텍스트와 표를 수집.

        본문이 그룹 도형 안에 들어있는 슬라이드가 많다 (실측: Lesson Learned
        본문 전체가 그룹). 최상위 도형만 훑으면 그 텍스트가 조용히 소실되므로
        그룹을 재귀로 펼친다.
        """
        texts: List[str] = []
        tables: List[List[List[str]]] = []
        PostbuyParser._collect_shapes(slide.shapes, texts, tables)
        return texts, tables

    @staticmethod
    def _collect_shapes(shapes, texts: List[str],
                        tables: List[List[List[str]]]) -> None:
        """도형 트리를 훑어 텍스트와 표를 모은다 (그룹 재귀)"""
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                PostbuyParser._collect_shapes(shape.shapes, texts, tables)
                continue
            if getattr(shape, 'has_table', False):
                rows = [[su.norm(c.text).replace('\n', ' ') for c in r.cells]
                        for r in shape.table.rows]
                if rows:
                    tables.append(rows)
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(re.sub(r'\s*\n\s*', ' / ', t))
        return texts, tables

    @staticmethod
    def _classify(texts: List[str]) -> str:
        blob = ' '.join(texts[:6]).lower()
        for kind, pattern in _SLIDE_KINDS:
            if re.search(pattern, blob, re.IGNORECASE):
                return kind
        return 'other'

    @staticmethod
    def _read_cover(texts: List[str], result: PostbuyResult) -> None:
        if texts:
            result.campaign = texts[0].replace('Post-buy Report', '').strip(' /')
        for t in texts[1:]:
            if re.search(r'20\d{2}', t):
                result.report_date = t.strip()
                break

    @staticmethod
    def _read_summary(section: PostbuySection, file_name: str,
                      result: PostbuyResult) -> None:
        """Campaign Summary — 총계 문구와 매체별 집행 결과 표"""
        for t in section.headlines:
            if '억' in t or '노출' in t:
                result.summary_headline = t
                break
        for t in section.headlines + section.notes:
            m = _PERIOD_RE.search(t)
            if m:
                result.period_raw = m.group(1).strip()
                break

        for ti, table in enumerate(section.tables, 1):
            if len(table) < 2:
                continue
            cols = PostbuyParser._table_columns(table[0])
            if 'impressions' not in cols:
                continue
            for ri, row in enumerate(table[1:], 1):
                media = row[0] if row else ''
                if not media or su.is_total_row(media):
                    # 🔴 Total 행을 그냥 버리지 않는다.
                    #
                    # 예전에는 여기서 건너뛰기만 해서, 보고서에서 가장 많이
                    # 인용되는 '매체비 총합'을 매체별 행을 더해 만들어야 했다.
                    # 병합셀로 매체명이 빈 행(카카오모먼트_크레딧 등)이 빠지면서
                    # 실측 캠페인에서 395,170,000 이어야 할 값이 414,170,000
                    # 으로 나왔다. 총계 행은 문서가 직접 말해 주는 값이므로
                    # 검증 기준으로 따로 보관한다 (claude.md 3.2 교차 인용).
                    if su.is_total_row(media):
                        budget = su.to_number(
                            PostbuyParser._at(row, cols.get('budget')))
                        if budget is not None and result.summary_total_budget is None:
                            result.summary_total_budget = budget
                            result.summary_total_source = (
                                f'{file_name}:p{section.slide_no}')
                    continue
                result.campaign_summary.append(MediaPerformance(
                    media=media,
                    product=PostbuyParser._at(row, cols.get('product')),
                    metrics=PostbuyParser._metrics_from(row, cols),
                    budget=su.to_number(PostbuyParser._at(row, cols.get('budget'))),
                    source=Provenance(file_name=file_name,
                                      sheet_or_slide=f'p{section.slide_no}',
                                      locator=f't{ti}r{ri}'),
                ))

    @staticmethod
    def _read_kpi(section: PostbuySection, file_name: str,
                  result: PostbuyResult) -> None:
        """
        KPI 달성 현황 — 두 가지 표가 온다.
          ① 구분 | Primary KPI | 목표 | 집행결과 | 달성율
          ② 구분 | 매체 | 상품 | KPI | 목표 | 집행결과 | 달성율
        헤더 표기가 병합으로 비어 있는 경우가 많아 위치 기반으로 읽는다.
        """
        for ti, table in enumerate(section.tables, 1):
            if len(table) < 2 or len(table[0]) < 4:
                continue

            body = [[su.norm(c) for c in row] for row in table[1:]]
            width = max((len(r) for r in body), default=0)
            if width < 4:
                continue

            # 라벨 열 판정: 그 열의 값이 대체로 숫자가 아니면 라벨 열
            label_cols: List[int] = []
            for j in range(width):
                vals = [r[j] for r in body if j < len(r) and r[j]]
                if not vals:
                    continue
                numeric = sum(1 for v in vals if su.to_number(v) is not None)
                if numeric <= len(vals) / 2:
                    label_cols.append(j)

            # 구분/매체는 병합으로 비어 있는 행이 많아 아래로 채운다
            filled = {j: su.forward_fill([r[j] if j < len(r) else '' for r in body])
                      for j in label_cols}

            for ri, cells in enumerate(body):
                nums = [su.to_number(c) for c in cells]
                nums = [v for v in nums if v is not None]
                if len(nums) < 3:
                    continue
                target, actual, rate = nums[-3:]
                labels = [filled[j][ri] for j in label_cols if filled[j][ri]]
                if not labels:
                    continue
                result.kpi_targets.append(KpiTarget(
                    scope=labels[0],
                    media=labels[1] if len(labels) > 2 else '',
                    product=labels[2] if len(labels) > 3 else '',
                    kpi_name=labels[-1],
                    target=target,
                    actual=actual,
                    achievement_rate=rate,
                    source=Provenance(file_name=file_name,
                                      sheet_or_slide=f'p{section.slide_no}',
                                      locator=f't{ti}r{ri + 1}'),
                ))

    @staticmethod
    def _read_events(section: PostbuySection, file_name: str,
                     result: PostbuyResult) -> None:
        """Scheduling 슬라이드의 날짜 주석"""
        for t in section.headlines + [h for h in section.title.split(' / ')]:
            m = _EVENT_DATE_RE.match(t)
            if not m:
                continue
            note = t[m.end():].strip(' /')
            if note:
                result.events.append(ScheduleEvent(
                    date=f'{int(m.group(1)):02d}-{int(m.group(2)):02d}',
                    note=note,
                    source=Provenance(file_name=file_name,
                                      sheet_or_slide=f'p{section.slide_no}'),
                ))

    # ------------------------------------------------------------------

    @staticmethod
    def _table_columns(header_row: List[str]) -> Dict[str, int]:
        cols: Dict[str, int] = {}
        for j, v in enumerate(header_row):
            canon = su.canonical_column(v)
            if canon and canon not in cols:
                cols[canon] = j
        return cols

    @staticmethod
    def _at(row: List[str], idx: Optional[int]) -> str:
        if idx is None or idx >= len(row):
            return ''
        return su.norm(row[idx])

    @staticmethod
    def _metrics_from(row: List[str], cols: Dict[str, int]) -> MetricSet:
        keys = ('billed', 'spend', 'impressions', 'views', 'clicks',
                'vtr', 'ctr', 'cpm', 'cpv', 'cpc')
        return MetricSet(**{
            k: su.to_number(PostbuyParser._at(row, cols.get(k))) for k in keys
        })
