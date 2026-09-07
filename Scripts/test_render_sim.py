# -*- coding: utf-8 -*-
"""
Stage 4 — PPTX 실제 렌더 시뮬레이션

    python Scripts/test_render_sim.py            # 실제 캠페인 6건 + 스트레스
    python Scripts/test_render_sim.py --stress   # 합성 스트레스만 (빠름)

여기까지가 이제껏 검증되지 않던 구간이다. 앞선 시뮬레이션은 **덱 구성**
(ReportSpec)까지만 봤다. 실제 파일을 만드는 단계는 전혀 다른 실패면을 갖는다.

  · 파일이 만들어졌는데 열리지 않는다
  · 슬라이드가 조용히 빠진다 (핸들러 누락 → 장수만 줄어듦)
  · 글자가 도형을 넘쳐 잘린다
  · 파일명 규칙(YYMMDD_품목_자료명_v0_Cheil)이 품목의 이상 문자로 깨진다
  · 큰 덱에서 시간·메모리가 튄다

'예외가 안 났다'를 합격으로 치지 않는다. 만들어진 파일을 **다시 열어서**
장수와 내용을 확인한다.
"""

import argparse
import io
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

from sim.runner import (CRITICAL, LOGIC_ISSUE, MINOR, PASS,     # noqa: E402
                        QUALITY_ISSUE, SILENT_FAILURE,
                        SYSTEM_ERROR, Result, render_table,
                        summarize, SEVERITY_ORDER)

SAMPLES = ROOT / 'Input' / 'Samples'

# 렌더 한 번이 이보다 오래 걸리면 '리포트 만들기'가 멈춘 것처럼 느껴진다
SLOW_RENDER_SEC = 25.0
# 무료 호스팅에서 다운로드가 부담스러워지는 크기
BIG_FILE_MB = 25.0


def _build(folder: Path):
    """캠페인 폴더 → (knowledge, dataset, spec)"""
    from utils.capture import capture_output
    from data_scanning import run_step1
    from utils.dataset_builder import CampaignDatasetBuilder
    from utils.report import ReportSpecBuilder

    log = io.StringIO()
    with capture_output(log):
        knowledge = run_step1(str(folder))
        dataset = CampaignDatasetBuilder.build(knowledge)
        spec = ReportSpecBuilder.build(knowledge, dataset)
    return knowledge, dataset, spec


def _render(spec, out_path: Path):
    """실제 PPTX 파일을 만든다. (경과초, 렌더러) 반환."""
    from utils.capture import capture_output
    from utils.report import ReportRenderer

    log = io.StringIO()
    r = ReportRenderer()
    t0 = time.time()
    with capture_output(log):
        r.render(spec, str(out_path))
    return time.time() - t0, r


def _inspect(path: Path) -> dict:
    """만들어진 파일을 **다시 열어** 확인한다."""
    from pptx import Presentation
    prs = Presentation(str(path))
    slides = list(prs.slides)
    empty = 0
    chars = 0
    for s in slides:
        text = ''
        for sh in s.shapes:
            text += (getattr(sh, 'text', '') or '')
            if getattr(sh, 'has_table', False):
                text += 'T'
            if sh.shape_type is not None and not getattr(sh, 'text', ''):
                text += ''
        chars += len(text)
        if not text.strip():
            empty += 1
    return {'slides': len(slides), 'empty': empty, 'chars': chars,
            'mb': path.stat().st_size / 1048576}


def check_deck(label: str, spec, out_dir: Path) -> Result:
    r = Result(scenario=label)
    out = out_dir / f'{label.replace("/", "_")[:40]}.pptx'
    try:
        secs, renderer = _render(spec, out)
    except Exception as e:
        r.fail(CRITICAL, f'렌더 중 예외 — {type(e).__name__}: {e}')
        return r
    r.seconds = secs

    if not out.is_file() or out.stat().st_size == 0:
        r.fail(CRITICAL, '파일이 만들어지지 않음')
        return r

    try:
        info = _inspect(out)
    except Exception as e:
        r.fail(CRITICAL, f'만든 파일을 다시 열 수 없음 — {type(e).__name__}: {e}')
        return r

    r.stats = {**info, 'spec_slides': len(spec.slides)}

    # 1) 조용히 빠진 슬라이드가 있는가
    skipped = getattr(renderer, 'skipped_kinds', []) or []
    if info['slides'] != len(spec.slides):
        lost = len(spec.slides) - info['slides']
        if skipped:
            r.fail(LOGIC_ISSUE,
                   f'슬라이드 {lost}장 누락 — 처리기 없는 종류: '
                   + ', '.join(sorted(set(skipped))))
        else:
            r.fail(SILENT_FAILURE,
                   f'슬라이드 {lost}장이 흔적 없이 사라짐 '
                   f'(구성 {len(spec.slides)} → 파일 {info["slides"]})')

    # 2) 빈 슬라이드
    if info['empty']:
        r.fail(QUALITY_ISSUE, f'내용이 없는 슬라이드 {info["empty"]}장')

    # 3) 성능
    if secs > SLOW_RENDER_SEC:
        r.fail(QUALITY_ISSUE, f'렌더에 {secs:.1f}초')
    if info['mb'] > BIG_FILE_MB:
        r.fail(QUALITY_ISSUE, f'파일이 {info["mb"]:.1f}MB — 내려받기 부담')

    # 4) 그리다 실패한 장 — 파일은 나왔지만 내용이 빈 장이 섞였다
    failed = getattr(renderer, 'failed_slides', []) or []
    if failed:
        kinds = ', '.join(sorted({k for k, _ in failed}))
        r.fail(LOGIC_ISSUE,
               f'슬라이드 {len(failed)}장을 그리지 못해 비워 둠 — {kinds} '
               f'({failed[0][1][:50]})')

    if skipped:
        r.notes.append(f'[INFO] 건너뛴 종류: {", ".join(sorted(set(skipped)))}')

    return r


# ═══════════════════════════ 실제 캠페인

def simulate_real(out_dir: Path) -> list:
    results = []
    if not SAMPLES.is_dir():
        return results
    for d in sorted(SAMPLES.iterdir()):
        if not d.is_dir() or d.name.startswith(('_', '.')):
            continue
        label = f'render/{d.name[:26]}'
        try:
            _, _, spec = _build(d)
        except Exception as e:
            r = Result(scenario=label)
            r.fail(CRITICAL, f'덱 구성 실패 — {type(e).__name__}: {e}')
            results.append(r)
            continue
        if not spec.slides:
            r = Result(scenario=label)
            r.fail(LOGIC_ISSUE, '슬라이드가 0장 구성됨')
            results.append(r)
            continue
        results.append(check_deck(label, spec, out_dir))
    return results


# ═══════════════════════════ 합성 스트레스

def _stress_specs():
    """렌더러를 일부러 괴롭히는 덱들."""
    from utils.report.blocks import ReportSpec, SlideSpec

    long_text = '매우 긴 문안 ' * 400
    emoji = '캠페인 🚀 성과 📈 ✨'
    yield ('stress/빈 덱', ReportSpec(campaign_tag='TAG', slides=[]))

    yield ('stress/아주 긴 문안', ReportSpec(
        campaign_tag='TAG',
        slides=[SlideSpec('lesson', {'section': long_text[:200],
                                     'bullets': [long_text] * 6})]))

    yield ('stress/이모지·특수문자', ReportSpec(
        campaign_tag=emoji,
        slides=[SlideSpec('divider', {'text': emoji}),
                SlideSpec('lesson', {'section': emoji,
                                     'bullets': [emoji, '<b>html</b>',
                                                 '{"json": 1}', 'a\tb\nc']})]))

    yield ('stress/빈 payload', ReportSpec(
        campaign_tag='TAG',
        slides=[SlideSpec('lesson', {}), SlideSpec('divider', {}),
                SlideSpec('toc', {}), SlideSpec('summary', {})]))

    yield ('stress/처리기 없는 종류', ReportSpec(
        campaign_tag='TAG',
        slides=[SlideSpec('divider', {'text': '정상'}),
                SlideSpec('nonexistent_kind', {'text': '없는 종류'})]))

    big_rows = [[f'매체{i}', f'{i*1000:,}', f'{i*99:,}', f'{i%100}%']
                for i in range(300)]
    yield ('stress/대용량 표', ReportSpec(
        campaign_tag='TAG',
        slides=[SlideSpec('media_table', {
            'section': '매체별 성과',
            'header': ['매체', '노출', '클릭', 'CTR'],
            'rows': big_rows})]))

    yield ('stress/많은 슬라이드(120장)', ReportSpec(
        campaign_tag='TAG',
        slides=[SlideSpec('divider', {'text': f'{i:03d}. 구간'})
                for i in range(120)]))


# 일부러 깨뜨려 넣은 시나리오 — 여기서 기대하는 건 '안 터지는 것'이 아니라
# **터지되 보고서 전체를 잃지 않는 것**이다. 강등이 확인되면 통과로 본다.
EXPECT_GRACEFUL = {'stress/빈 payload', 'stress/처리기 없는 종류'}


def simulate_stress(out_dir: Path) -> list:
    results = []
    for label, spec in _stress_specs():
        if not spec.slides:
            # 빈 덱은 '파일이 생기되 0장'이 정상 — 별도로 본다
            r = Result(scenario=label)
            out = out_dir / 'empty.pptx'
            try:
                secs, _ = _render(spec, out)
                r.seconds = secs
                if not out.is_file():
                    r.fail(SYSTEM_ERROR, '빈 덱에서 파일이 만들어지지 않음')
            except Exception as e:
                r.fail(CRITICAL, f'빈 덱 렌더에서 예외 — {type(e).__name__}: {e}')
            results.append(r)
            continue
        r = check_deck(label, spec, out_dir)
        if label in EXPECT_GRACEFUL and r.grade in (LOGIC_ISSUE, QUALITY_ISSUE):
            # 파일이 만들어졌고 나머지 장이 살아 있으면 의도한 강등이다
            r.notes.insert(0, '[기대] 실패한 장만 비워 두고 덱은 살아남음')
            r.grade = PASS
        results.append(r)
    return results


# ═══════════════════════════ 파일명 규칙

def simulate_filenames() -> list:
    """
    출력 파일명 규칙(YYMMDD_품목_자료명_v0_Cheil)이 이상한 품목에서도
    실제로 저장 가능한 이름을 만드는지 (claude.md 1.4).
    """
    results = []
    hostile = [
        ('정상', '냉장고'),
        ('슬래시', 'AV/VD'),
        ('역슬래시', 'A\\B'),
        ('콜론', '에어컨:무풍'),
        ('물음표', '뭐지?'),
        ('별표', 'TV*'),
        ('따옴표', '"큰따옴표"'),
        ('이모지', '냉장고🧊'),
        ('아주 긺', '가' * 200),
        ('빈값', ''),
        ('공백만', '   '),
        ('점으로 끝', '품목.'),
        ('예약어', 'CON'),
    ]
    from utils.errors import safe_filename_part
    tmp = Path(tempfile.mkdtemp(prefix='axname_'))
    try:
        for label, cat in hostile:
            r = Result(scenario=f'filename/{label}')
            # 제품이 쓰는 것과 같은 정제 함수를 통과시킨다.
            # 정제 없이 그대로 쓰면 '/', '?', '*' 에서 저장이 실패한다.
            safe = safe_filename_part(cat, fallback='품목')
            name = f'260908_{safe}_결과리포트_v0_Cheil.pptx'
            try:
                p = tmp / name
                p.write_bytes(b'x')
                if not p.is_file():
                    r.fail(SYSTEM_ERROR, '파일이 생성되지 않음')
                elif not safe:
                    r.fail(LOGIC_ISSUE, '품목이 빈 문자열로 정제됨')
            except OSError as e:
                r.fail(SYSTEM_ERROR,
                       f'정제 후에도 저장 불가 — {type(e).__name__}: '
                       f'{str(e)[:50]}')
            results.append(r)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return results


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--stress', action='store_true', help='합성 스트레스만')
    args = ap.parse_args()

    out_dir = Path(tempfile.mkdtemp(prefix='axrender_'))
    started = time.time()
    try:
        results = simulate_stress(out_dir) + simulate_filenames()
        if not args.stress:
            results += simulate_real(out_dir)
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)

    print('\n' + '=' * 78)
    print('Stage 4 렌더 시뮬레이션')
    print('=' * 78)
    print(render_table(results))

    counts = summarize(results)
    total = len(results)
    parts = [f'{g} {counts[g]}' for g in SEVERITY_ORDER if counts[g]]
    print('\n' + '-' * 78)
    print(f'전체 {total}건 · ' + ' · '.join(parts))
    print(f'소요 {time.time() - started:.1f}초')
    bad = total - counts[PASS] - counts[MINOR]
    print(f'조치 필요 {bad}건')
    return 1 if bad else 0


if __name__ == '__main__':
    raise SystemExit(main())
